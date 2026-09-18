"""
Document parser — PDF, DOCX, TXT, images.

Returns a Document object (veritas.models.document.Document).
"""
from __future__ import annotations

from pathlib import Path
from typing import Optional

from veritas.models.document import Document, Paragraph
from veritas.tools.text_utils import (
    clean_text,
    count_words,
    split_paragraphs,
    is_heading,
    detect_references,
)
from veritas.tools.ocr import (
    is_ocr_available,
    ocr_image,
    ocr_pdf_page,
)


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".png", ".jpg", ".jpeg"}


class DocumentParser:
    """Parse documents into Document objects."""

    def __init__(self, ocr_threshold: int = 50, verbose: bool = True):
        """
        Args:
            ocr_threshold: if a PDF page has fewer than this many chars
                           of extractable text, OCR it.
            verbose: print progress.
        """
        self.ocr_threshold = ocr_threshold
        self.verbose = verbose

    # ---------- Public ----------
    def parse(self, path: str | Path) -> Document:
        """Parse any supported file into a Document."""
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {path}")

        ext = path.suffix.lower()
        if ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(f"Unsupported file type: {ext}")

        if self.verbose:
            print(f"[parser] Parsing {path.name} ({ext})")

        if ext == ".pdf":
            doc = self._parse_pdf(path)
        elif ext == ".docx":
            doc = self._parse_docx(path)
        elif ext in (".txt", ".md"):
            doc = self._parse_txt(path)
        elif ext in (".png", ".jpg", ".jpeg"):
            doc = self._parse_image(path)
        else:
            raise ValueError(f"Unhandled extension: {ext}")

        # Post-process: detect references, split paragraphs
        doc = self._finalize(doc, path)
        return doc

    # ---------- PDF ----------
    def _parse_pdf(self, path: Path) -> Document:
        try:
            import pymupdf  # noqa
            doc = pymupdf.open(str(path))
        except ImportError:
            import fitz  # type: ignore
            doc = fitz.open(str(path))

        all_pages_text = []
        page_count = doc.page_count
        ocr_used_pages = []

        for page_num in range(page_count):
            page = doc[page_num]
            text = page.get_text("text") or ""

            if len(text.strip()) < self.ocr_threshold:
                # Scanned page — try OCR
                if is_ocr_available():
                    if self.verbose:
                        print(f"[parser]   OCR page {page_num + 1}")
                    try:
                        text = ocr_pdf_page(page)
                        ocr_used_pages.append(page_num + 1)
                    except Exception as e:
                        if self.verbose:
                            print(f"[parser]   OCR failed: {e}")
                        text = ""

            all_pages_text.append(text)

        doc.close()

        full_text = clean_text("\n\n".join(all_pages_text))
        return Document(
            path=str(path),
            title=path.stem,
            text=full_text,
            page_count=page_count,
            metadata={
                "source_type": "pdf",
                "ocr_pages": ocr_used_pages,
            },
        )

    # ---------- DOCX ----------
    def _parse_docx(self, path: Path) -> Document:
        from docx import Document as DocxDocument

        docx_doc = DocxDocument(str(path))
        paragraphs_text = [p.text for p in docx_doc.paragraphs if p.text.strip()]
        full_text = clean_text("\n\n".join(paragraphs_text))

        return Document(
            path=str(path),
            title=path.stem,
            text=full_text,
            page_count=1,  # DOCX doesn't have fixed pages
            metadata={"source_type": "docx"},
        )

    # ---------- TXT ----------
    def _parse_txt(self, path: Path) -> Document:
        # Try common encodings
        for enc in ("utf-8", "utf-16", "latin-1"):
            try:
                text = path.read_text(encoding=enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = path.read_bytes().decode("utf-8", errors="ignore")

        text = clean_text(text)
        return Document(
            path=str(path),
            title=path.stem,
            text=text,
            page_count=1,
            metadata={"source_type": "txt"},
        )

    # ---------- Image ----------
    def _parse_image(self, path: Path) -> Document:
        if not is_ocr_available():
            raise RuntimeError(
                "OCR unavailable. Install Tesseract and pytesseract."
            )
        text = ocr_image(path)
        text = clean_text(text)
        return Document(
            path=str(path),
            title=path.stem,
            text=text,
            page_count=1,
            metadata={"source_type": "image", "ocr": True},
        )

    # ---------- Post-processing ----------
    def _finalize(self, doc: Document, path: Path) -> Document:
        """Split into paragraphs, detect headings, extract references."""
        paragraphs = split_paragraphs(doc.text, min_chars=30)

        ref_lines, body_text = detect_references(doc.text)

        para_objs = []
        for i, text in enumerate(paragraphs):
            # Is this in the references section?
            is_ref = bool(
                ref_lines and text in "\n".join(ref_lines)
            )
            para_objs.append(Paragraph(
                text=text,
                page=0,
                index=i,
                is_reference=is_ref,
                is_heading=is_heading(text),
                word_count=count_words(text),
            ))

        doc.paragraphs = para_objs
        doc.references = ref_lines
        doc.word_count = count_words(doc.text)
        doc.metadata["body_text"] = body_text

        if self.verbose:
            print(
                f"[parser] Done: {doc.word_count} words, "
                f"{len(para_objs)} paragraphs, "
                f"{len(ref_lines)} references"
            )
        return doc


def parse_document(path: str | Path, verbose: bool = True) -> Document:
    """Convenience function."""
    return DocumentParser(verbose=verbose).parse(path)


__all__ = ["DocumentParser", "parse_document", "SUPPORTED_EXTENSIONS"]